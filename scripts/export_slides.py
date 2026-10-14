"""Export the rendered deck to PowerPoint and to a PDF of slides with speaker notes.

Reads docs/slides.html (run `quarto render` first) in headless Chromium and writes:
  exports/imperial-careers-slides.pptx        one picture per slide state, speaker notes attached;
                                              animated slides and the opening atlas video embedded
                                              as MP4s that start when the slide opens
  exports/imperial-careers-slides-notes.pdf   one page per slide: the slide, then its notes

Slides with fragments become one slide per state (before and after the reveal). Animated slides
are captured from Chrome's screencast at full resolution and encoded with ffmpeg.

    python3 scripts/export_slides.py
"""
import asyncio
import base64
import html
import shutil
import subprocess
import tempfile
import time
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from threading import Thread

from lxml import etree
from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'exports'
W, H = 1920, 1080
ATLAS_VIDEO = Path.home() / 'col_matching/docs/combined_mobility.mp4'
ATLAS_POSTER = ROOT / 'images/combined-mobility-frame.jpg'
HIDE = '.motion-replay,.motion-tag{display:none!important}'


class Quiet(SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass


async def capture(base, work):
    """Return one dict per exported slide: image, notes, and an optional video."""
    async with async_playwright() as p:
        browser = await p.chromium.launch(args=['--no-sandbox'])
        ctx = await browser.new_context(viewport={'width': W, 'height': H}, reduced_motion='reduce')
        page = await ctx.new_page()
        await page.goto(base + '/slides.html', wait_until='networkidle')
        await page.wait_for_function('window.Reveal?.isReady()')
        await page.add_style_tag(content=HIDE)
        await page.evaluate("Reveal.configure({transition: 'none', backgroundTransition: 'none', controls: false})")
        info = await page.evaluate("""() => Reveal.getSlides().map(s => ({
            id: s.id, motion: s.classList.contains('motion'),
            video: !!s.querySelector('video'),
            fragments: s.querySelectorAll('.fragment').length,
            notes: (() => {
                const a = s.querySelector('aside.notes');
                if (!a) return '';
                const c = a.cloneNode(true);
                c.querySelectorAll('style, script, mjx-assistive-mml').forEach(e => e.remove());
                const paras = [...c.querySelectorAll('p, li')].map(p => p.textContent.trim()).filter(Boolean);
                return (paras.length ? paras.join('\\n') : c.textContent).trim();
            })(),
            title: (() => {
                const h = (s.querySelector('h2') || s.querySelector('h1'))?.cloneNode(true);
                if (!h) return '';
                h.querySelectorAll('.motion-tag').forEach(e => e.remove());
                return h.textContent.trim();
            })()}))""")
        slides = []
        for i, s in enumerate(info):
            for f in range(-1, s['fragments']):
                await page.evaluate('([i, f]) => Reveal.slide(i, 0, f)', [i, f])
                await page.wait_for_timeout(700)
                img = work / f'slide-{i:02d}-{f + 1}.png'
                await page.screenshot(path=str(img))
                slides.append({'index': i, 'image': img, 'notes': s['notes'], 'title': s['title'],
                               'motion': s['motion'] and f == s['fragments'] - 1, 'id': s['id'],
                               'video': s['video']})
        await ctx.close()

        for s in [s for s in slides if s['motion']]:
            s['movie'] = await record(browser, base, s['index'], work)
        await browser.close()
    return slides


async def record(browser, base, index, work):
    """Screencast one animated slide from a blank first frame to a held final frame."""
    ctx = await browser.new_context(viewport={'width': W, 'height': H})
    page = await ctx.new_page()
    await page.goto(base + '/slides.html?motion', wait_until='networkidle')
    await page.wait_for_function('window.Reveal?.isReady()')
    await page.add_style_tag(content=HIDE)
    await page.evaluate("Reveal.configure({transition: 'none', controls: false})")
    await page.evaluate('i => Reveal.slide(i)', index)
    last = await page.evaluate("Math.max(...[...Reveal.getCurrentSlide().querySelectorAll('[data-at]')]"
                               ".map(e => +e.dataset.at))")
    await page.wait_for_timeout(last + 2000)
    frames = []
    cdp = await ctx.new_cdp_session(page)

    def on_frame(event):
        frames.append((event['metadata']['timestamp'], event['data']))
        asyncio.ensure_future(cdp.send('Page.screencastFrameAck', {'sessionId': event['sessionId']}))

    cdp.on('Page.screencastFrame', on_frame)
    await cdp.send('Page.startScreencast', {'format': 'jpeg', 'quality': 92, 'maxWidth': W, 'maxHeight': H})
    await page.wait_for_timeout(400)
    started = time.time()
    await page.evaluate("Reveal.getCurrentSlide().querySelector('.motion-replay').click()")
    hold = 2.5
    await page.wait_for_timeout(last + 1800)
    await cdp.send('Page.stopScreencast')
    await ctx.close()

    frames = [f for f in sorted(frames) if f[0] >= started - 0.02]
    folder = work / f'frames-{index:02d}'
    folder.mkdir()
    lines = []
    for n, (stamp, data) in enumerate(frames):
        path = folder / f'{n:05d}.jpg'
        path.write_bytes(base64.b64decode(data))
        nxt = frames[n + 1][0] if n + 1 < len(frames) else stamp + hold
        lines += [f"file '{path}'", f'duration {max(nxt - stamp, 0.001):.4f}']
    lines.append(f"file '{folder / f'{len(frames) - 1:05d}.jpg'}'")
    (folder / 'list.txt').write_text('\n'.join(lines) + '\n')
    movie = work / f'motion-{index:02d}.mp4'
    subprocess.run(['ffmpeg', '-v', 'error', '-y', '-f', 'concat', '-safe', '0', '-i', str(folder / 'list.txt'),
                    '-vf', f'scale={W}:{H},fps=30', '-c:v', 'libx264', '-pix_fmt', 'yuv420p', '-crf', '20',
                    '-an', str(movie)], check=True)
    return movie


P = 'http://schemas.openxmlformats.org/presentationml/2006/main'


def autoplay(slide, shape_id, seconds):
    """Replace python-pptx's click-to-play timing with play-when-the-slide-opens."""
    old = slide._element.find(f'{{{P}}}timing')
    if old is not None:
        slide._element.remove(old)
    xml = f'''<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{int(seconds * 1000)}" fill="hold"/><p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''
    timing = etree.fromstring(xml)
    ext = slide._element.find(f'{{{P}}}extLst')
    if ext is not None:
        ext.addprevious(timing)
    else:
        slide._element.append(timing)


def duration(movie):
    out = subprocess.run(['ffprobe', '-v', 'error', '-show_entries', 'format=duration', '-of', 'csv=p=0',
                          str(movie)], capture_output=True, text=True, check=True)
    return float(out.stdout.strip())


def build_pptx(slides, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for s in slides:
        slide = prs.slides.add_slide(blank)
        movie, poster = s.get('movie'), s['image']
        if s['video'] and ATLAS_VIDEO.exists():
            movie, poster = ATLAS_VIDEO, ATLAS_POSTER
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(0x15, 0x16, 0x1a)
        if movie:
            shape = slide.shapes.add_movie(str(movie), 0, 0, prs.slide_width, prs.slide_height,
                                           poster_frame_image=str(poster), mime_type='video/mp4')
            autoplay(slide, shape.shape_id, duration(movie))
        else:
            slide.shapes.add_picture(str(s['image']), 0, 0, prs.slide_width, prs.slide_height)
        if s['notes']:
            slide.notes_slide.notes_text_frame.text = s['notes']
    prs.save(path)


async def build_pdf(slides, base_dir, path):
    pages = []
    for n, s in enumerate(slides, 1):
        notes = ''.join(f'<p>{html.escape(par)}</p>' for par in s['notes'].split('\n') if par.strip())
        pages.append(f'''<section><header>Slide {n}{' · ' + html.escape(s['title']) if s['title'] else ''}</header>
<img src="{s['image'].name}">{notes or '<p class="none">No speaker notes.</p>'}</section>''')
    doc = f'''<!doctype html><meta charset="utf-8"><style>
@page {{ size: letter; margin: 0.55in; }}
body {{ font: 10.5pt/1.45 Georgia, serif; color: #222; margin: 0; }}
section {{ page-break-after: always; }}
header {{ font: 600 9pt system-ui, sans-serif; color: #666; letter-spacing: .03em; margin-bottom: .4em; }}
img {{ width: 100%; border: 1px solid #ccc; margin-bottom: .8em; }}
p {{ margin: 0 0 .6em; }} .none {{ color: #888; font-style: italic; }}
</style>
<h1 style="font: 400 20pt Georgia, serif; page-break-after: always; margin-top: 3in">Cleaning, Structuring and Grounding Historical Data<br>
<span style="font-size: 12pt; color: #555">Slides with speaker notes · AI and History Conference, Johns Hopkins, 15 October 2026 · Jim Clifford</span></h1>
{''.join(pages)}'''
    (base_dir / 'notes.html').write_text(doc)
    async with async_playwright() as p:
        browser = await p.chromium.launch(args=['--no-sandbox'])
        page = await browser.new_page()
        await page.goto((base_dir / 'notes.html').as_uri(), wait_until='load')
        await page.pdf(path=str(path), format='Letter', print_background=True)
        await browser.close()


async def main():
    if not shutil.which('ffmpeg'):
        raise SystemExit('ffmpeg is required')
    OUT.mkdir(exist_ok=True)
    server = ThreadingHTTPServer(('127.0.0.1', 0), partial(Quiet, directory=str(ROOT / 'docs')))
    Thread(target=server.serve_forever, daemon=True).start()
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        slides = await capture(f'http://127.0.0.1:{server.server_port}', work)
        pptx_path = OUT / 'imperial-careers-slides.pptx'
        pdf_path = OUT / 'imperial-careers-slides-notes.pdf'
        build_pptx(slides, pptx_path)
        await build_pdf(slides, work, pdf_path)
    movies = sum(1 for s in slides if s.get('movie') or s['video'])
    print(f'{len(slides)} slides ({movies} with video) -> {pptx_path} ({pptx_path.stat().st_size / 1e6:.1f} MB)')
    print(f'notes PDF -> {pdf_path} ({pdf_path.stat().st_size / 1e6:.1f} MB)')


if __name__ == '__main__':
    asyncio.run(main())
